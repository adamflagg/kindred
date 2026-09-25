"""Local parity check: the financial-aid calculator against an exported aid sheet.

LOCAL ONLY. It never runs in CI and no committed test gives it real data. Its
inputs -- a sheet export, a rules document and a parity config -- live in the
gitignored docs/plans/campership-data/ of the main checkout. This file holds the
sheet's column positions and the comparison logic, and no data.

It runs calculate() over every sheet row and counts, per output column, how
many rows match the sheet (campership design section 8):

  sheet    grants exactly as the sheet applied them (Aid Calculator column P).
           Must match every row: this is the acceptance gate.
  correct  grants joined by CampMinder person id from the Grants tab, cancelled
           grants left out. Differences are expected; review them against the
           known-defect list.

It prints counts and sheet ROW NUMBERS only -- never a name, an id or an amount.

    uv run python -m scripts.financial_aid.parity_check \\
        --rules  <main checkout>/docs/plans/campership-data/aid-rules-2026.json \\
        --config <main checkout>/docs/plans/campership-data/parity-config-2026.json \\
        --sheet  <main checkout>/docs/plans/campership-data/<export>.xlsx \\
        --mode both

Run it as a module from the repository root (`-m`): run by file path, Python puts
scripts/financial_aid/ on sys.path instead of the root, and `bunking` fails to
import `api`.

Exit 0: sheet mode matched every row (or only correct mode ran). Exit 1: a
mismatch, a rules-validation error, or a catalog session missing from the map.
Exit 2: the export's layout is not the one this file expects.
"""

from __future__ import annotations

import argparse
import hashlib
import sys
from collections import defaultdict
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal

import openpyxl
from pydantic import BaseModel, ConfigDict, Field, field_validator

from bunking.financial_aid.calculator import (
    ApplicationInputs,
    CalcResult,
    CostOverride,
    GrantInput,
    RequestInputs,
    calculate,
)
from bunking.financial_aid.rules import AidRules, resolve_program, validate_rules

Mode = Literal["sheet", "correct"]

RAW_DATA = "Raw Data"
AID_CALCULATOR = "Aid Calculator"
REFERENCES = "References"
GRANTS = "Grants"

# key -> (column letter, header text lowercased and stripped). A header that does
# not match stops the run, so a moved column cannot feed the wrong figure.
RAW_COLUMNS: dict[str, tuple[str, str]] = {
    "unique_id": ("A", "unique id"),
    "personal_id": ("G", "personal id"),
    "py_gross": ("N", "prior year gross pre-tax income:"),
    "py_confirm": ("O", "confirm py gross pre-tax income"),
    "py_agi": ("P", "prior year adjusted gross income:"),
    "cy_gross": ("Q", "current year expected gross pre-tax income"),
    "savings": ("S", "non-retirement savings & investments"),
    "medical": ("U", "expected medical expenses"),
    "education": ("V", "expected education/student loan expenses"),
    "unemployment": ("AA", "fa-unemployment"),
    "single_parent": ("AB", "fa-singleparent"),
    "bipoc": ("AF", "does your child identify as a person of color?"),
    "dependents": ("AH", "number of dependents"),
    "gender_identity": ("AJ", "gender identity"),
    "pronouns": ("AK", "pronouns"),
}
CALC_COLUMNS: dict[str, tuple[str, str]] = {
    "unique_id": ("A", "unique id"),
    "stage": ("C", "stage"),
    "session": ("F", "session(s)"),
    "ask": ("G", "aid requested"),
    "override": ("I", "family camp or override costs"),
    "total": ("J", "total award"),
    "income": ("K", "updated calculated income"),
    "income_tier": ("L", "income tier"),
    "final_tier": ("M", "final tier"),
    "potential": ("N", "potential grant"),
    "grants": ("P", "grants"),
    "r1": ("Q", "round 1 grant"),
    "appeal": ("U", "r2 aid requested:"),
    "r2": ("W", "round 2 grant"),
    "discretionary": ("X", "discretionary funds granted"),
}
REF_COLUMNS: dict[str, tuple[str, str]] = {
    "session": ("A", "session"),
    "cost": ("B", "cost"),
    "program_type": ("C", "program type"),
}
GRANT_COLUMNS: dict[str, tuple[str, str]] = {
    "cancelled": ("A", "cancelled?"),
    "camper_id": ("E", "camper id"),
    "amount": ("K", "grant amount"),
}

COMPARED: tuple[tuple[str, str, Callable[[CalcResult], Decimal | int | None]], ...] = (
    ("income", "adjusted income (K)", lambda r: r.adjusted_income),
    ("income_tier", "income tier (L)", lambda r: r.income_tier),
    ("final_tier", "final tier (M)", lambda r: r.final_tier),
    ("potential", "R1 potential (N)", lambda r: r.r1_potential),
    ("grants", "grants (P)", lambda r: r.grants_offset),
    ("r1", "Round 1 (Q)", lambda r: r.r1),
    ("r2", "Round 2 (W)", lambda r: r.r2),
    ("total", "total (J)", lambda r: r.total),
)
# The sheet computes in floating point; anything within this is the same number.
TOLERANCE = Decimal("0.000001")
_MODE_LABELS = {"sheet": "sheet-grant", "correct": "correct-grant"}
_RAW_NUMBERS = ("py_gross", "py_confirm", "py_agi", "cy_gross", "savings", "medical", "education")


class SheetLayoutError(ValueError):
    """The export's tabs or headers are not where this harness reads them."""


class ParityConfig(BaseModel):
    """Local, gitignored: how the sheet's text maps onto CampMinder ids and program keys."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    # References session text -> the 2026 CampMinder session it names.
    session_map: dict[str, int]
    # References program type text -> the rules document's program key (a cross-check).
    program_type_map: dict[str, str]
    # Stage label -> decision type key (e.g. the full-cost stage).
    stage_decision_types: dict[str, str] = Field(default_factory=dict)
    # Program key for rows whose session text is not in References.
    unmatched_program: str
    # Reason code recorded for the sheet's typed cost column I.
    override_reason: str
    first_row: int = 2
    last_row: int | None = None
    # The sheet's catalog range ends at row 44 (SessionInfo); rows below it were never read.
    catalog_last_row: int = 44

    @field_validator("session_map", "program_type_map", "stage_decision_types")
    @classmethod
    def _lowercase_keys(cls, value: dict[str, Any]) -> dict[str, Any]:
        # The sheet's text comparisons are case-insensitive.
        return {key.lower(): item for key, item in value.items()}


@dataclass(frozen=True)
class SheetRow:
    row: int
    raw: dict[str, Any]
    calc: dict[str, Any]


@dataclass(frozen=True)
class Sheet:
    rows: list[SheetRow]
    catalog: dict[str, tuple[Decimal | None, str]]
    grants_by_person: dict[str, list[Decimal]]


@dataclass
class Diagnostics:
    unmapped_catalog_sessions: set[str] = field(default_factory=set)
    routing_mismatch_rows: set[int] = field(default_factory=set)
    price_conflicts: set[str] = field(default_factory=set)
    # Unparsable: a non-blank cell that could not be read as a number.
    unparsable: defaultdict[str, int] = field(default_factory=lambda: defaultdict(int))
    # Blank: a required field (only `ask` today) whose cell was empty, so it fell
    # back to 0 rather than the sheet's own figure. Never set for an optional
    # field -- a blank override, appeal or discretionary amount is legitimate.
    blank: defaultdict[str, int] = field(default_factory=lambda: defaultdict(int))


@dataclass(frozen=True)
class ModeReport:
    mode: Mode
    rows: int
    matched: dict[str, int]
    mismatched_rows: dict[str, list[int]]

    @property
    def passed(self) -> bool:
        return all(count == self.rows for count in self.matched.values())


def _number(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int | float):
        return Decimal(str(value))
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("$", "")
        if not text:
            return None
        try:
            return Decimal(text)
        except InvalidOperation:
            return None
    return None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    return value if isinstance(value, str) else str(value)


def _id_text(value: Any) -> str | None:
    number = _number(value)
    if number is not None and number == number.to_integral_value():
        return str(int(number))
    text = _text(value)
    return text.strip() if text else None


def _cell(ws: Any, column: str, row: int) -> Any:
    return ws[f"{column}{row}"].value


def _read(ws: Any, columns: dict[str, tuple[str, str]], row: int) -> dict[str, Any]:
    return {key: _cell(ws, column, row) for key, (column, _) in columns.items()}


def _check_headers(ws: Any, tab: str, columns: dict[str, tuple[str, str]]) -> None:
    for key, (column, expected) in columns.items():
        raw_value = _cell(ws, column, 1)
        found = str(raw_value or "").strip().lower()
        if found != expected:
            # Never echo what the cell held -- a moved header row could carry a
            # name, amount or other private value. Only the shape of the miss.
            raise SheetLayoutError(
                f"{tab}!{column}1 does not read {expected!r} ({key}); "
                f"found a {type(raw_value).__name__} of length {len(found)}"
            )


def load_sheet(path: Path, config: ParityConfig) -> Sheet:
    workbook = openpyxl.load_workbook(path, data_only=True)
    tabs = {}
    for tab, columns in (
        (RAW_DATA, RAW_COLUMNS),
        (AID_CALCULATOR, CALC_COLUMNS),
        (REFERENCES, REF_COLUMNS),
        (GRANTS, GRANT_COLUMNS),
    ):
        if tab not in workbook.sheetnames:
            raise SheetLayoutError(f"the export has no {tab!r} tab")
        tabs[tab] = workbook[tab]
        _check_headers(tabs[tab], tab, columns)

    raw_by_id: dict[str, dict[str, Any]] = {}
    raw_row_of: dict[str, int] = {}
    raw_ws = tabs[RAW_DATA]
    for row in range(2, raw_ws.max_row + 1):
        values = _read(raw_ws, RAW_COLUMNS, row)
        unique_id = _text(values["unique_id"])
        if not unique_id:
            continue
        if unique_id in raw_row_of:
            # Never name the id itself -- only the two row numbers that collide.
            raise SheetLayoutError(f"{RAW_DATA} rows {raw_row_of[unique_id]} and {row} share a Unique ID")
        raw_by_id[unique_id] = values
        raw_row_of[unique_id] = row

    # Rows join on the sheet's Unique ID, never on row position.
    rows: list[SheetRow] = []
    calc_ws = tabs[AID_CALCULATOR]
    for row in range(config.first_row, (config.last_row or calc_ws.max_row) + 1):
        calc = _read(calc_ws, CALC_COLUMNS, row)
        unique_id = _text(calc["unique_id"])
        if not unique_id:
            continue
        raw = raw_by_id.get(unique_id)
        if raw is None:
            raise SheetLayoutError(f"{AID_CALCULATOR} row {row} has no {RAW_DATA} row with the same Unique ID")
        rows.append(SheetRow(row=row, raw=raw, calc=calc))

    catalog: dict[str, tuple[Decimal | None, str]] = {}
    for row in range(2, config.catalog_last_row + 1):
        values = _read(tabs[REFERENCES], REF_COLUMNS, row)
        text = _text(values["session"])
        if text is None:
            continue
        key = text.lower()
        if key not in catalog:  # the first catalog row wins, as the sheet's lookup did
            catalog[key] = (_number(values["cost"]), (_text(values["program_type"]) or "").lower())

    grants: defaultdict[str, list[Decimal]] = defaultdict(list)
    grant_ws = tabs[GRANTS]
    for row in range(2, grant_ws.max_row + 1):
        values = _read(grant_ws, GRANT_COLUMNS, row)
        person = _id_text(values["camper_id"])
        amount = _number(values["amount"])
        cancelled = (_text(values["cancelled"]) or "").strip().lower() == "yes"
        if person and amount is not None and amount > 0 and not cancelled:
            grants[person].append(amount)
    return Sheet(rows=rows, catalog=catalog, grants_by_person=dict(grants))


def _counted_number(source: dict[str, Any], key: str, diagnostics: Diagnostics) -> Decimal | None:
    """A number read from any sheet dict (raw or calc). A non-blank cell that
    fails to parse is counted; a blank cell is not -- callers decide whether
    blank is legitimate (most fields) or must itself be counted (`_required_number`)."""
    value = source[key]
    number = _number(value)
    if number is None and value not in (None, ""):
        diagnostics.unparsable[key] += 1
    return number


def _required_number(source: dict[str, Any], key: str, diagnostics: Diagnostics) -> Decimal:
    """A required Money field (only `ask` today). Blank and unparsable are both
    counted, in separate buckets, and both fall back to 0 -- the calculator needs
    *a* value either way, but principle 5 says that fallback must never be silent."""
    value = source[key]
    number = _number(value)
    if number is not None:
        return number
    if value in (None, "") or (isinstance(value, str) and not value.strip()):
        diagnostics.blank[key] += 1
    else:
        diagnostics.unparsable[key] += 1
    return Decimal(0)


def _dependents(raw: dict[str, Any], diagnostics: Diagnostics) -> int | None:
    number = _counted_number(raw, "dependents", diagnostics)
    if number is None:
        return None
    if number != number.to_integral_value() or number < 0:
        diagnostics.unparsable["dependents"] += 1
        return None
    return int(number)


def _grants(row: SheetRow, sheet: Sheet, mode: Mode) -> list[GrantInput]:
    if mode == "sheet":
        applied = _number(row.calc["grants"])
        return [GrantInput(amount=applied, state="committed")] if applied is not None and applied > 0 else []
    person = _id_text(row.raw["personal_id"])
    if not person:
        return []
    return [GrantInput(amount=amount, state="committed") for amount in sheet.grants_by_person.get(person, [])]


@dataclass(frozen=True)
class ParsedRow:
    """Everything about one row that does NOT depend on --mode. Computed once per
    row so a blank or unparsable cell is counted once, not once per mode."""

    application: ApplicationInputs
    session_cm_id: int | None
    program_key: str
    ask: Decimal
    equity_answers: dict[str, Any]
    cost_override: CostOverride | None
    appeal_amount: Decimal | None
    decision_type: str | None
    discretionary_amount: Decimal


def _parse_row(
    row: SheetRow, sheet: Sheet, rules: AidRules, config: ParityConfig, diagnostics: Diagnostics
) -> ParsedRow:
    raw, calc = row.raw, row.calc
    numbers = {key: _counted_number(raw, key, diagnostics) for key in _RAW_NUMBERS}
    application = ApplicationInputs(
        prior_year_gross=numbers["py_gross"],
        prior_year_confirmed=numbers["py_confirm"],
        prior_year_agi=numbers["py_agi"],
        current_year_gross=numbers["cy_gross"],
        savings=numbers["savings"],
        medical_expenses=numbers["medical"],
        education_expenses=numbers["education"],
        dependents=_dependents(raw, diagnostics),
        answers={"unemployment": _text(raw["unemployment"]), "single_parent": _text(raw["single_parent"])},
    )

    session_text = (_text(calc["session"]) or "").lower()
    session_cm_id: int | None = None
    program_key = config.unmatched_program
    if session_text in sheet.catalog:
        mapped = config.session_map.get(session_text)
        if mapped is None:
            diagnostics.unmapped_catalog_sessions.add(session_text)
        else:
            session_cm_id = mapped
            resolved = resolve_program(rules, mapped)
            price, program_type = sheet.catalog[session_text]
            if resolved is None or resolved != config.program_type_map.get(program_type):
                diagnostics.routing_mismatch_rows.add(row.row)
            if price != rules.cost.tuition.get(mapped):
                diagnostics.price_conflicts.add(session_text)
            program_key = resolved or config.unmatched_program

    stage = (_text(calc["stage"]) or "").lower()
    decision_type = config.stage_decision_types.get(stage)
    decision = rules.awards.decision_types.get(decision_type) if decision_type else None
    discretionary = _counted_number(calc, "discretionary", diagnostics) or Decimal(0)
    if decision is not None and decision.kind == "full_cost":
        # On these rows the sheet's X is the top-up the calculator computes itself.
        discretionary = Decimal(0)
    override = _counted_number(calc, "override", diagnostics)
    # ask is a required Money on RequestInputs: a blank cell still needs *a*
    # value to build the input, but that fallback to 0 is counted, never silent.
    ask = _required_number(calc, "ask", diagnostics)
    appeal_amount = _counted_number(calc, "appeal", diagnostics)
    return ParsedRow(
        application=application,
        session_cm_id=session_cm_id,
        program_key=program_key,
        ask=ask,
        equity_answers={
            "bipoc": _text(raw["bipoc"]),
            "gender_identity": _text(raw["gender_identity"]),
            "pronouns": _text(raw["pronouns"]),
        },
        cost_override=CostOverride(amount=override, reason=config.override_reason) if override is not None else None,
        appeal_amount=appeal_amount,
        decision_type=decision_type,
        discretionary_amount=discretionary,
    )


def _build_request(parsed: ParsedRow, row: SheetRow, sheet: Sheet, mode: Mode) -> RequestInputs:
    # Only the grant source differs by mode; everything else was parsed once.
    return RequestInputs(
        session_cm_id=parsed.session_cm_id,
        program_key=parsed.program_key,
        ask=parsed.ask,
        equity_answers=parsed.equity_answers,
        cost_override=parsed.cost_override,
        grants_applicable=_grants(row, sheet, mode),
        appeal_amount=parsed.appeal_amount,
        decision_type=parsed.decision_type,
        discretionary_amount=parsed.discretionary_amount,
    )


def _matches(sheet_value: Any, ours: Decimal | int | None, *, blank_is_zero: bool) -> bool:
    if isinstance(sheet_value, str) and sheet_value.startswith("#"):
        return ours is None  # a sheet error matches only a result the calculator could not compute
    expected = _number(sheet_value)
    if expected is None:
        if not blank_is_zero:
            return ours is None
        expected = Decimal(0)
    if ours is None:
        return False
    return abs(expected - Decimal(ours)) <= TOLERANCE


def run_mode(sheet: Sheet, parsed_rows: dict[int, ParsedRow], rules: AidRules, mode: Mode) -> ModeReport:
    matched = {key: 0 for key, _, _ in COMPARED}
    mismatched: dict[str, list[int]] = {key: [] for key, _, _ in COMPARED}
    for row in sheet.rows:
        parsed = parsed_rows[row.row]
        request = _build_request(parsed, row, sheet, mode)
        result = calculate(parsed.application, request, rules)
        for key, _, ours in COMPARED:
            if _matches(row.calc[key], ours(result), blank_is_zero=key == "grants"):
                matched[key] += 1
            else:
                mismatched[key].append(row.row)
    return ModeReport(mode=mode, rows=len(sheet.rows), matched=matched, mismatched_rows=mismatched)


def _describe(path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()[:12]
    modified = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).strftime("%Y-%m-%d %H:%M UTC")
    return f"{path.name} sha256={digest} modified={modified}"


def _print_report(report: ModeReport) -> None:
    print(f"mode {_MODE_LABELS[report.mode]}: rows={report.rows}")
    for key, label, _ in COMPARED:
        print(f"  {label:<22} {report.matched[key]}/{report.rows}")
        rows = report.mismatched_rows[key]
        if rows:
            more = " ..." if len(rows) > 25 else ""
            print(f"    mismatched sheet rows: {', '.join(str(r) for r in rows[:25])}{more}")


def _print_diagnostics(diagnostics: Diagnostics) -> None:
    print("diagnostics:")
    print(f"  catalog sessions missing from session_map: {sorted(diagnostics.unmapped_catalog_sessions) or 'none'}")
    print(f"  rows routed to a program other than the sheet's: {sorted(diagnostics.routing_mismatch_rows) or 'none'}")
    print(f"  catalog prices that differ from the rules' tuition: {sorted(diagnostics.price_conflicts) or 'none'}")
    print(f"  unparsable cells by column: {dict(diagnostics.unparsable) or 'none'}")
    print(f"  blank required cells by column: {dict(diagnostics.blank) or 'none'}")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Local parity check of the financial-aid calculator.")
    parser.add_argument("--rules", type=Path, required=True)
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--sheet", type=Path, required=True)
    parser.add_argument("--mode", choices=("sheet", "correct", "both"), default="both")
    args = parser.parse_args(argv)

    rules = AidRules.model_validate_json(args.rules.read_text(encoding="utf-8"))
    config = ParityConfig.model_validate_json(args.config.read_text(encoding="utf-8"))
    print(f"oracle: {_describe(args.sheet)}")
    report = validate_rules(rules)
    print(f"rules:  {args.rules.name} year={rules.year} errors={len(report.errors)} warnings={len(report.warnings)}")
    for issue in report.issues:
        print(f"  {issue.severity:<7} {issue.code} at {issue.path}")
    if report.errors:
        print("RESULT: FAIL (the rules document has validation errors)")
        return 1
    if config.unmatched_program not in rules.programs:
        print(f"RESULT: FAIL (unmatched_program {config.unmatched_program!r} is not a program in the rules)")
        return 1
    try:
        sheet = load_sheet(args.sheet, config)
    except SheetLayoutError as error:
        print(f"RESULT: LAYOUT ({error})")
        return 2

    diagnostics = Diagnostics()
    # Parsed once per row, independent of --mode: --mode both must not double-count
    # a blank or unparsable cell just because it runs the row twice.
    parsed_rows = {row.row: _parse_row(row, sheet, rules, config, diagnostics) for row in sheet.rows}

    modes: list[Mode] = ["sheet", "correct"] if args.mode == "both" else [args.mode]
    reports = [run_mode(sheet, parsed_rows, rules, mode) for mode in modes]
    for mode_report in reports:
        _print_report(mode_report)
    _print_diagnostics(diagnostics)

    sheet_report = next((r for r in reports if r.mode == "sheet"), None)
    unmapped = bool(diagnostics.unmapped_catalog_sessions)
    failed = unmapped or (sheet_report is not None and not sheet_report.passed)
    if sheet_report is not None:
        verdict = "FAIL" if failed else "PASS"
    elif unmapped:
        # --mode correct alone still fails on an unmapped catalog session, but
        # says why instead of printing "n/a" over an exit code of 1.
        verdict = "FAIL (catalog sessions missing from session_map; sheet mode did not run)"
    else:
        verdict = "n/a"
    print(f"RESULT sheet-grant: {verdict}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
