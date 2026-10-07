"""Load one season's arrival curve from its applications workbook (Scenarios addendum §S11.7; owner §S15 items 1, 3).

AGGREGATES ONLY. It reads the workbook's "Raw Data" sheet, finds the column headed "Timestamp" (by its header, never
its position), skips fully blank rows (trailing template rows), and turns each timestamp into its camp day. It then
builds the weekly cumulative curve and prints only counts, the first and last day and the weekly shares. No cell's
text, name or id is ever printed or stored, and an unreadable row stops the load naming its row NUMBER only.

Timestamps: a date-time cell, or text shaped "YYYY/MM/DD h:mm:ss am TZ" (TZ PST, PDT or PT) or "M/D/YYYY h:mm:ss am".
All are Pacific wall-clock time.

    uv run python -m scripts.financial_aid.load_arrival_curve --workbook <path> --year 2026 [--deadline 2026-02-04] [--write]

The anchor is --deadline, else the year's approved milestones.application_deadline (read through the superuser
client), else Jan 1 (stored as "calendar"). Without --write nothing is written; with it the season's row in
aid_arrival_curves is replaced. It is run once for 2026 by the owner's prod agent, and again for a prior year only if
that year's workbook turns up. It never runs in CD and ships no data; the workbook never enters any repository.

Run it as a module from the repository root (`-m`), as parity_check is.
"""

from __future__ import annotations

import argparse
import asyncio
import os
import sys
from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Final

import openpyxl

from api.services.financial_aid_arrival_curves_repository import ArrivalCurveRepository
from api.services.financial_aid_rules_service import AidRulesRepository, FinancialAidRulesService
from bunking.financial_aid.arrival import ArrivalCurve, calendar_anchor, camp_date_of, curve_from_dates
from pocketbase import PocketBase
from scripts.utils.auth import authenticate_pocketbase

RAW_DATA: Final = "Raw Data"
HEADER: Final = "timestamp"
LOADER: Final = "system:arrival-curve-loader"
_PACIFIC: Final = frozenset({"PST", "PDT", "PT"})
_MERIDIEM: Final = frozenset({"AM", "PM"})  # a trailing "AM" is the clock's, not a zone
_TENTH: Final = Decimal("0.1")


class WorkbookError(Exception):
    """The workbook can't be read as asked. The message names a sheet, a header or a row number: never a cell."""


def _strptime(text: str, shape: str) -> datetime | None:
    try:
        return datetime.strptime(text, shape)
    except ValueError:
        return None


def parse_timestamp(cell: object) -> datetime | None:
    """A date-time cell, or one of the two text shapes the 2026 sheet holds; None for anything else. The year-first
    shape ends in a zone ("2026/03/08 8:30:00 am PDT"); the month-first shape has none, and its am/pm may be in either
    case ("3/8/2026 9:00:00 am" or "… AM")."""
    if isinstance(cell, datetime):
        return cell
    if not isinstance(cell, str):
        return None
    text = " ".join(cell.split())
    head, _, zone = text.rpartition(" ")
    if head and zone.isalpha() and zone.isupper() and zone not in _MERIDIEM:
        return _strptime(head, "%Y/%m/%d %I:%M:%S %p") if zone in _PACIFIC else None
    return _strptime(text, "%m/%d/%Y %I:%M:%S %p")


def _blank(row: Sequence[object]) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def read_timestamps(path: Path) -> list[datetime]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if RAW_DATA not in book.sheetnames:
            raise WorkbookError(f'No sheet named "{RAW_DATA}"')
        rows = book[RAW_DATA].iter_rows(values_only=True)
        header = next(rows, ()) or ()
        columns = [i for i, cell in enumerate(header) if isinstance(cell, str) and cell.strip().lower() == HEADER]
        if len(columns) != 1:
            raise WorkbookError(f'{RAW_DATA} row 1 has {"no" if not columns else "more than one"} "Timestamp" header')
        (column,) = columns
        moments: list[datetime] = []
        for number, row in enumerate(rows, start=2):
            if _blank(row):
                continue
            moment = parse_timestamp(row[column] if column < len(row) else None)
            if moment is None:
                raise WorkbookError(f"{RAW_DATA} row {number}: the Timestamp is not a date and time")
            moments.append(moment)
        if not moments:
            raise WorkbookError(f"{RAW_DATA} has no timestamped rows")
        return moments
    finally:
        book.close()


def summary_lines(curve: ArrivalCurve, days: Sequence[date]) -> list[str]:
    aligned = (
        f"the application deadline ({curve.anchor.isoformat()})"
        if curve.aligned_on == "application_deadline"
        else f"the calendar (from {curve.anchor.isoformat()})"
    )
    lines = [
        f"{curve.year} arrival curve: aligned on {aligned}",
        f"counted {curve.counted} · first {min(days).isoformat()} · last {max(days).isoformat()}",
    ]
    lines += [f"week {point.week}: {(point.share * 100).quantize(_TENTH)}%" for point in curve.points]
    return lines


def _client() -> PocketBase:  # authenticated from the environment
    return authenticate_pocketbase(os.getenv("POCKETBASE_URL", "http://localhost:8090"))


async def approved_deadline(year: int) -> date | None:
    approved = await FinancialAidRulesService(AidRulesRepository(_client(), read_only=True)).latest_approved(
        year, ["milestones"]
    )
    return approved.document.milestones.application_deadline if approved is not None else None


async def save_curve(curve: ArrivalCurve) -> None:
    await ArrivalCurveRepository(_client()).save(curve, actor=LOADER)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Load one season's arrival curve (aggregates only).")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--deadline", type=date.fromisoformat, default=None)
    parser.add_argument("--write", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    if args.write and not os.getenv("POCKETBASE_URL"):
        # authenticate_pocketbase defaults to the local dev server and dev credentials: never write there silently.
        print("Set POCKETBASE_URL to the PocketBase to write to. Nothing was written.", file=sys.stderr)
        return 2
    try:
        moments = read_timestamps(args.workbook)
    except WorkbookError as exc:
        print(f"{exc}. Nothing was written.", file=sys.stderr)
        return 2
    days = [camp_date_of(moment) for moment in moments]
    deadline = args.deadline if args.deadline is not None else asyncio.run(approved_deadline(args.year))
    curve = curve_from_dates(
        days,
        deadline if deadline is not None else calendar_anchor(args.year),
        year=args.year,
        source="workbook",
        aligned_on="application_deadline" if deadline is not None else "calendar",
    )
    for line in summary_lines(curve, days):
        print(line)
    if not args.write:
        print("Dry run: nothing written (add --write to store it).")
        return 0
    asyncio.run(save_curve(curve))
    print(f"Stored {args.year}'s arrival curve.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
