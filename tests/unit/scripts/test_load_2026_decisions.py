"""The 2026 decision-year load (D67 as amended by D145 and the owner's 10-07 rulings). Fictional only: the rules are
the engine fixture's invented season, every id is in the 1000xxx range, and the one workbook is built here.

Fixture arithmetic (tier 2, Session 1000102 costs 4,000): Round 1 is 75% = 3,000; with an appeal of 1,000 Round 2
is 90% − 3,000 = 600. The full-cost program tops Round 1 up to cost + 50 = 4,050.
"""

from __future__ import annotations

import json
from collections.abc import Mapping, Sequence
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from api.services.financial_aid_reconciliation import CampLine
from bunking.financial_aid.calculator import GrantInput
from bunking.financial_aid.change_log import AidWrite
from bunking.financial_aid.decisions import DecisionEvent, fold_rounds, price_request, season_budget
from bunking.financial_aid.decisions.pricing import RequestToPrice
from bunking.financial_aid.rules import AidRules
from scripts.financial_aid import load_2026_decisions as loader
from scripts.financial_aid.load_2026_decisions import LOADER, LoadRequest, SheetAward, plan_load, plan_writes
from scripts.financial_aid.parity_check import CALC_COLUMNS, GRANT_COLUMNS, RAW_COLUMNS, REF_COLUMNS, ParityConfig
from tests.unit.bunking.financial_aid.fixtures import app, fictional_rules, req, with_lever

EMMA = "reqemma00000001"
LIAM = "reqliam00000001"
FAMILY = "reqfamily000001"
T0 = datetime(2031, 9, 1, tzinfo=UTC)


def _rules(**levers: Any) -> AidRules:
    rules = fictional_rules()
    for path, value in levers.items():
        rules = with_lever(rules, path.replace("__", "."), value)
    return rules


RULES = _rules()
# The full-cost type outside the budget, as 2026's outside-funded program is (D121).
OUTSIDE = with_lever(RULES, "awards.decision_types.full_cost_program.counts_toward_budget", False)


def _award(row: int = 2, **fields: Any) -> SheetAward:
    base: dict[str, Any] = {
        "row": row,
        "person_cm_id": 1000002,
        "household_cm_id": 1000001,
        "session_cm_id": 1000102,
        "stage": "1 Accepted",
        "application": app(),
        "request": req(),
        "ask": Decimal(4000),
        "appeal": None,
        "r1": Decimal(3000),
        "r2": None,
        "extra": Decimal(0),
        "decision_type": None,
    }
    base.update(fields)
    return SheetAward(**base)


def _request(
    rid: str = EMMA, *, household: int = 1000001, person: int = 1000002, session: int = 1000102, status: str = "active"
) -> LoadRequest:
    return LoadRequest(id=rid, household_cm_id=household, person_cm_id=person, session_cm_id=session, status=status)


def _line(
    amount: str, *, txn: int = 1000900, household: int = 1000001, person: int = 1000002, session: int = 1000102
) -> CampLine:
    return CampLine(
        transaction_cm_id=txn,
        household_cm_id=household,
        person_cm_id=person,
        amount=Decimal(amount),
        post_date=T0,
        is_reversed=False,
        reversal_date=None,
        attributed_person_cm_id=person,
        attributed_session_cm_id=session,
    )


def _plan(
    awards: Sequence[SheetAward],
    requests: Sequence[LoadRequest] = (),
    lines: Sequence[CampLine] = (),
    existing: Sequence[DecisionEvent] = (),
    rules: AidRules = RULES,
    trackers: dict[int, str] | None = None,
    overrides: Sequence[loader.LoadOverride] = (),
) -> loader.LoadPlan:
    return plan_load(
        awards,
        requests or [_request()],
        lines,
        existing,
        rules,
        rules_version=3,
        trackers=trackers or {},
        overrides=overrides,
    )


def _posts(plan: loader.LoadPlan, rid: str = EMMA) -> dict[int, Decimal]:
    return {
        e.round: e.amount for e in plan.creates if e.request_id == rid and e.kind == "post" and e.amount is not None
    }


def _snap(event: loader.PlannedEvent) -> Mapping[str, Any]:
    assert event.snapshot is not None
    return event.snapshot


def _as_events(plan: loader.LoadPlan) -> list[DecisionEvent]:
    """The planned rows as the app would read them back, in the order they would be created."""
    return [
        DecisionEvent(
            id=f"ev{i:013d}",
            request_id=e.request_id,
            round=e.round,
            kind=e.kind,
            created=T0,
            amount=e.amount,
            decision_type=e.decision_type,
            lock_source=e.lock_source,
            rules_version=e.rules_version,
            snapshot=e.snapshot,
            note=e.note,
            actor=LOADER,
        )
        for i, e in enumerate(plan.creates)
    ]


def _kinds(plan: loader.LoadPlan) -> dict[str, int]:
    out: dict[str, int] = {}
    for row in plan.report:
        out[row.kind] = out.get(row.kind, 0) + 1
    return out


# --- what a matched request is loaded as ------------------------------------------------------------------------------


def test_a_matched_request_posts_its_round_at_campminders_net_marked_reproduced_by_the_loader() -> None:
    plan = _plan([_award()], lines=[_line("3000")])
    (post,) = [e for e in plan.creates if e.kind == "post"]
    assert (post.request_id, post.round, post.amount) == (EMMA, 1, Decimal(3000))
    assert (post.lock_source, post.rules_version) == ("reproduced", 3)
    assert post.effective_on is None  # 2026 has no decision dates (D67)
    assert _snap(post)["decided"] == "3000"  # the engine's receipt, reproduced from the sheet's inputs
    assert _snap(post)["result"]["r1"] == "3000"
    assert _snap(post)["reproduced"]["sheet"] == "3000"
    assert _snap(post)["reproduced"]["campminder"] == "3000"
    assert all(e.data(2031)["actor"] == LOADER for e in plan.creates)
    # "1 Accepted": the family accepted the posted round.
    assert [(e.kind, e.round) for e in plan.creates if e.kind == "accept"] == [("accept", 1)]
    assert plan.report == []  # nothing to resolve: sheet, engine and CampMinder agree


def test_an_offered_stage_posts_without_accepting() -> None:
    plan = _plan([_award(stage="1 Offered")], lines=[_line("3000")])
    assert [e.kind for e in plan.creates] == ["post"]


def test_campminder_wins_and_the_difference_lands_on_the_latest_posted_round() -> None:
    award = _award(stage="2 Accepted", appeal=Decimal(1000), r2=Decimal(600))
    plan = _plan([award], lines=[_line("3000", txn=1000901), _line("400", txn=1000902)])  # CampMinder holds 3,400
    assert _posts(plan) == {1: Decimal(3000), 2: Decimal(400)}  # 3,600 on the sheet: −200 on Round 2
    r2 = next(e for e in plan.creates if e.kind == "post" and e.round == 2)
    assert _snap(r2)["decided"] == "600"  # the sheet's round, kept beside CampMinder's money
    asks = [e for e in plan.creates if e.kind == "ask"]
    assert [(e.round, e.amount) for e in asks] == [(2, Decimal(1000))]
    (row,) = plan.report
    assert (row.kind, row.sheet_amount, row.campminder_amount, row.posted_amount) == (
        "campminder_differs",
        Decimal(3600),
        Decimal(3400),
        Decimal(3400),
    )
    assert row.difference == Decimal(-200)


def test_a_difference_larger_than_the_latest_round_spills_back_to_the_round_before() -> None:
    award = _award(stage="2 Accepted", appeal=Decimal(1000), r2=Decimal(600))
    plan = _plan([award], lines=[_line("2500")])
    assert _posts(plan) == {1: Decimal(2500), 2: Decimal(0)}


def test_campminder_above_the_sheet_raises_the_latest_round() -> None:
    plan = _plan([_award()], lines=[_line("3250")], trackers={1000001: "E"})
    assert _posts(plan) == {1: Decimal(3250)}
    (row,) = plan.report
    assert (row.kind, row.tracker, row.difference) == ("campminder_differs", "E", Decimal(250))


def test_a_household_campminder_holds_no_money_for_posts_zero() -> None:
    plan = _plan([_award()])
    assert _posts(plan) == {1: Decimal(0)}
    (row,) = plan.report
    assert (row.kind, row.campminder_amount, row.posted_amount) == ("campminder_differs", Decimal(0), Decimal(0))


def test_money_campminder_could_not_place_keeps_the_sheets_amount_and_is_reported_as_q_l12() -> None:
    # Two requests in two sessions and one household-level line naming neither: it can't be split.
    requests = [_request(), _request(LIAM, person=1000003, session=1000101)]
    awards = [
        _award(),
        _award(
            row=3,
            person_cm_id=1000003,
            session_cm_id=1000101,
            request=req(person_cm_id=1000003, session_cm_id=1000101),
            r1=Decimal(1500),
        ),
    ]
    line = _line("5000", person=0, session=0)
    plan = _plan(awards, requests, [line])
    assert _posts(plan, EMMA) == {1: Decimal(3000)}
    assert _posts(plan, LIAM) == {1: Decimal(1500)}
    unplaced = [r for r in plan.report if r.kind == "campminder_unplaced"]
    assert [(r.tracker, r.campminder_amount) for r in unplaced] == [("Q-L12", Decimal(5000))]
    assert _kinds(plan) == {"campminder_unplaced": 1, "campminder_unmatched": 2}


def test_siblings_in_one_session_are_matched_as_the_family_and_session_together() -> None:
    # Owner 10-07: CampMinder's money per family x session. A household-level line on the session covers both.
    requests = [_request(), _request(LIAM, person=1000003)]
    awards = [_award(), _award(row=3, person_cm_id=1000003, request=req(person_cm_id=1000003))]
    plan = _plan(awards, requests, [_line("6000", person=0)])
    assert (_posts(plan, EMMA), _posts(plan, LIAM)) == ({1: Decimal(3000)}, {1: Decimal(3000)})
    assert plan.report == []
    short = _plan(awards, requests, [_line("5900", person=0)])
    assert sorted([*_posts(short, EMMA).values(), *_posts(short, LIAM).values()]) == [Decimal(2900), Decimal(3000)]
    (row,) = short.report
    assert (row.kind, row.sheet_amount, row.campminder_amount, row.person_cm_id) == (
        "campminder_differs",
        Decimal(6000),
        Decimal(5900),
        None,  # the family x session, not one camper
    )


def test_a_line_on_a_session_the_household_has_no_request_for_is_reported() -> None:
    plan = _plan([_award()], lines=[_line("3000"), _line("700", txn=1000903, session=1000104)])
    assert _posts(plan) == {1: Decimal(3000)}
    assert [(r.kind, r.campminder_amount, r.session_cm_id) for r in plan.report] == [
        ("campminder_without_load", Decimal(700), 1000104)
    ]


def test_a_persons_request_in_another_household_does_not_take_this_households_line() -> None:
    elsewhere = _request(LIAM, household=1000005)  # the same camper on a second household's application
    plan = _plan([_award()], [_request(), elsewhere], [_line("3000")])
    assert _posts(plan) == {1: Decimal(3000)}
    assert plan.report == []


def test_a_household_line_with_one_request_in_the_household_is_that_requests() -> None:
    plan = _plan([_award()], lines=[_line("3000", person=0, session=0)])
    assert _posts(plan) == {1: Decimal(3000)}
    assert plan.report == []


def test_the_family_camp_request_is_found_by_household_and_session() -> None:
    family = _request(FAMILY, household=1000001, person=0, session=1000201)
    award = _award(
        session_cm_id=1000201,
        request=req(person_cm_id=1000002, session_cm_id=1000201, program_key="family_camp", headcount={"standard": 3}),
        r1=Decimal(1350),
    )
    plan = _plan([award], [family], [_line("1350", person=0, session=1000201)])
    assert _posts(plan, FAMILY) == {1: Decimal(1350)}


# --- named decision types and extra money ------------------------------------------------------------------------------


def test_a_full_cost_stage_keys_its_decision_type_and_its_round_sits_outside_the_budget() -> None:
    award = _award(stage="Full-cost program", extra=Decimal(1050), decision_type="full_cost_program")
    plan = _plan([award], lines=[_line("4050")], rules=OUTSIDE)
    (keyed,) = [e for e in plan.creates if e.kind == "award"]
    assert (keyed.round, keyed.decision_type, keyed.amount) == (1, "full_cost_program", Decimal(0))
    (post,) = [e for e in plan.creates if e.kind == "post"]
    assert post.amount == Decimal(4050)
    assert _snap(post)["counts_toward_budget"] is False
    assert {r.tracker for r in plan.report} <= {"outside"}


def test_a_difference_on_a_request_outside_the_budget_is_tracked_as_outside_by_no_program_name() -> None:
    award = _award(stage="Full-cost program", extra=Decimal(1050), decision_type="full_cost_program")
    plan = _plan([award], lines=[_line("4000")], rules=OUTSIDE)
    (row,) = [r for r in plan.report if r.kind == "campminder_differs"]
    assert row.tracker == "outside"


def test_a_300_stage_keys_the_rules_top_up_on_round_2() -> None:
    award = _award(stage="2 Accepted + $300", appeal=Decimal(1000), r2=Decimal(600), extra=Decimal(250))
    plan = _plan([award], lines=[_line("3850")])
    (keyed,) = [e for e in plan.creates if e.kind == "award"]
    assert (keyed.round, keyed.decision_type, keyed.amount) == (2, "appeal_top_up", Decimal(0))
    assert _posts(plan) == {1: Decimal(3000), 2: Decimal(850)}
    assert plan.report == []


def test_other_extra_money_is_the_rules_discretionary_type_on_round_3() -> None:
    award = _award(stage="3 Accepted", appeal=Decimal(1000), r2=Decimal(600), extra=Decimal(700))
    plan = _plan([award], lines=[_line("4300")])
    (keyed,) = [e for e in plan.creates if e.kind == "award"]
    assert (keyed.round, keyed.decision_type, keyed.amount) == (3, "discretionary", Decimal(700))
    assert _posts(plan) == {1: Decimal(3000), 2: Decimal(600), 3: Decimal(700)}
    assert {(e.kind, e.round) for e in plan.creates if e.kind == "accept"} == {
        ("accept", 1),
        ("accept", 2),
        ("accept", 3),
    }


def test_a_sheet_figure_the_engine_does_not_reproduce_is_reported_and_loaded_as_the_engine_has_it() -> None:
    plan = _plan([_award(r1=Decimal(2900))], lines=[_line("3000")])
    # Both differences are reported: the engine's against the sheet, and the posted money against the sheet.
    assert _kinds(plan) == {"engine_differs": 1, "campminder_differs": 1}
    row = next(r for r in plan.report if r.kind == "engine_differs")
    assert (row.sheet_amount, row.engine_amount) == (Decimal(2900), Decimal(3000))


def test_a_hold_staff_decided_past_is_released_so_the_engine_still_prices_it() -> None:
    # The placeholder-income check holds an income under 1,000; 2026's staff decided the row anyway.
    award = _award(application=app(prior_year_gross="500", current_year_gross="500"), r1=Decimal(3600))
    plan = _plan([award], lines=[_line("3600")])
    assert _posts(plan) == {1: Decimal(3600)}
    assert "placeholder_income" in _snap(plan.creates[0])["reproduced"]["released_holds"]


def test_the_sheets_grants_reach_the_engine() -> None:
    award = _award(request=req(grants_applicable=[GrantInput(amount=Decimal(500), state="committed")]))
    plan = _plan([award], lines=[_line("2500")])
    assert _snap(plan.creates[0])["decided"] == "2500"
    # The award said 3,000 before the grant: reported, never hidden.
    assert _kinds(plan) == {"engine_differs": 1, "campminder_differs": 1}


# --- what is never invented -----------------------------------------------------------------------------------------


def test_rows_that_match_no_request_are_reported_and_never_loaded() -> None:
    awards = [
        _award(row=2, person_cm_id=None),  # a weekend row with no CampMinder id
        _award(row=3, session_cm_id=None),  # the session text matched nothing
        _award(row=4, person_cm_id=1000002, session_cm_id=1000101),  # the person's request is another session
        _award(row=5, person_cm_id=1000009, household_cm_id=1000008),  # no request at all
    ]
    plan = _plan(awards, lines=[_line("3000")])
    assert plan.creates == []
    assert _kinds(plan) == {
        "no_campminder_id": 1,
        "session_unmatched": 1,
        "session_mismatch": 1,
        "no_request": 1,
        "campminder_without_load": 1,  # Emma's money sits on a request the sheet doesn't load
    }
    assert {r.tracker for r in plan.report} == {"unmatched"}


def test_two_sheet_rows_on_one_request_are_both_reported_and_neither_is_loaded() -> None:
    plan = _plan([_award(row=2), _award(row=3)], lines=[_line("3000")])
    assert [e for e in plan.creates if e.request_id == EMMA] == []
    assert [r.sheet_row for r in plan.report if r.kind == "duplicate"] == [2, 3]


# --- requests that are no longer live (coordinator ruling 10-07: skip and report, never post onto them) -------------


@pytest.mark.parametrize("status", ["withdrawn", "duplicate"])
def test_a_sheet_row_on_a_withdrawn_or_duplicate_request_is_reported_and_never_posted_onto(status: str) -> None:
    plan = _plan([_award()], [_request(status=status)], [_line("3000")])
    assert plan.creates == []
    (row,) = [r for r in plan.report if r.kind == "request_not_live"]
    assert (row.sheet_row, row.person_cm_id) == (2, 1000002)
    assert status in row.note


def test_an_unmatched_session_request_is_still_live_and_loads() -> None:
    plan = _plan([_award()], [_request(status="unmatched_session")], [_line("3000")])
    assert _posts(plan) == {1: Decimal(3000)}


@pytest.mark.parametrize("status", ["withdrawn", "duplicate", "duplicate_pending"])
def test_a_sheet_row_loads_onto_the_live_request_beside_its_closed_twin(status: str) -> None:
    # Intake keeps a duplicate in its holder's family x person x session slot, so the live one is the match.
    twin = _request(LIAM, status=status)
    plan = _plan([_award()], [twin, _request()], [_line("3000")])
    assert _posts(plan) == {1: Decimal(3000)}
    assert _posts(plan, LIAM) == {}
    assert "request_ambiguous" not in _kinds(plan)


def test_a_session_override_onto_a_closed_request_is_not_listed_as_loaded() -> None:
    rule = _override("session", person=1000002, sheet_session=1000101, session=1000102)
    plan = _plan([_award(session_cm_id=1000101)], [_request(status="withdrawn")], [_line("3000")], overrides=[rule])
    assert plan.creates == []
    assert "request_not_live" in _kinds(plan)
    assert "override_applied" not in _kinds(plan)  # the row lists only what loads under an override (owner, 10-07)


# --- the owner's overrides: each group's answer as data, never code (owner rulings pending 10-07) -------------------


def _override(decision: str, **fields: Any) -> loader.LoadOverride:
    return loader.LoadOverride(decision=decision, **fields)


def test_without_an_override_a_session_mismatch_stays_unloaded() -> None:
    plan = _plan([_award(session_cm_id=1000101)], lines=[_line("3000")])
    assert plan.creates == []
    assert "session_mismatch" in _kinds(plan)


def test_a_session_override_loads_the_sheet_row_onto_the_persons_request_in_that_session() -> None:
    award = _award(session_cm_id=1000101)  # the sheet says Session B; the request and CampMinder say Session A
    rule = _override("session", person=1000002, sheet_session=1000101, session=1000102, note="follows CampMinder")
    plan = _plan([award], lines=[_line("3000")], overrides=[rule])
    assert _posts(plan) == {1: Decimal(3000)}
    post = next(e for e in plan.creates if e.kind == "post")
    assert _snap(post)["reproduced"]["override"] == {"decision": "session", "sheet_session": 1000101}
    assert "session_mismatch" not in _kinds(plan)
    # Each row loaded under an override is listed, in the override's own words (owner, 10-07).
    (row,) = [r for r in plan.report if r.kind == "override_applied"]
    assert (row.sheet_row, row.session_cm_id, row.note) == (2, 1000102, "follows CampMinder")


def test_a_session_override_onto_a_request_with_no_session_reads_campminder_by_the_sheets_session() -> None:
    # An unmatched intake (no session on the request): the sheet's session is the only session evidence.
    request = _request(session=0, status="unmatched_session")
    award = _award(session_cm_id=1000101)
    rule = _override("session", person=1000002, sheet_session=1000101, session=0)
    plan = _plan([award], [request], [_line("3200", session=1000101)], overrides=[rule])
    assert _posts(plan) == {1: Decimal(3200)}
    assert "campminder_without_load" not in _kinds(plan)
    assert plan.campminder_on_loaded == Decimal(3200)


def test_a_session_override_is_held_to_its_household_and_sheet_row_when_it_names_them() -> None:
    award = _award(session_cm_id=1000101)
    elsewhere = _override("session", person=1000002, household=1000099, sheet_session=1000101, session=1000102)
    other_row = _override("session", person=1000002, sheet_session=1000101, sheet_row=9, session=1000102)
    plan = _plan([award], lines=[_line("3000")], overrides=[elsewhere, other_row])
    assert plan.creates == []
    assert _kinds(plan)["override_unused"] == 2


def test_a_skip_override_leaves_one_of_two_rows_on_a_request_out_and_loads_the_other() -> None:
    awards = [_award(row=2), _award(row=3, r1=Decimal(500))]
    rule = _override("skip", person=1000002, sheet_session=1000102, sheet_row=3)
    plan = _plan(awards, lines=[_line("3000")], overrides=[rule])
    assert _posts(plan) == {1: Decimal(3000)}
    assert "duplicate" not in _kinds(plan)
    (skipped,) = [r for r in plan.report if r.kind == "override_skip"]
    assert skipped.sheet_row == 3


def test_a_person_override_names_the_campminder_person_a_weekend_row_lacks() -> None:
    award = _award(row=4, person_cm_id=None, household_cm_id=None)
    rule = _override("person", sheet_row=4, sheet_session=1000102, to_person=1000002, to_household=1000001)
    plan = _plan([award], lines=[_line("3000")], overrides=[rule])
    assert _posts(plan) == {1: Decimal(3000)}
    assert "no_campminder_id" not in _kinds(plan)


def test_a_campminder_override_records_campminders_money_as_round_1_on_a_request_with_no_sheet_row() -> None:
    rule = _override("campminder", person=1000002, household=1000001, session=1000102)
    plan = _plan([], lines=[_line("2500")], overrides=[rule], trackers={1000001: "C"})
    (post,) = plan.creates  # posted only: with no sheet row there is no stage, so nothing is accepted
    assert post.note.startswith("CampMinder only")  # labelled as the owner asked (10-07)
    assert (post.request_id, post.round, post.kind, post.amount) == (EMMA, 1, "post", Decimal(2500))
    assert (post.lock_source, post.rules_version) == ("reproduced", 3)
    assert _snap(post)["reproduced"] == {
        "sheet_row": None,
        "stage": "",
        "sheet": None,
        "engine": None,
        "released_holds": [],
        "campminder": "2500",
        "override": {"decision": "campminder"},
    }
    (row,) = plan.report
    assert (row.kind, row.tracker, row.posted_amount, row.campminder_amount) == (
        "campminder_only",
        "C",
        Decimal(2500),
        Decimal(2500),
    )
    assert plan.campminder_on_loaded == Decimal(2500)
    assert plan.campminder_unloaded == Decimal(0)
    again = _plan([], lines=[_line("2500")], existing=_as_events(plan), overrides=[rule])
    assert again.unchanged


def test_a_campminder_override_on_a_request_campminder_holds_nothing_for_records_zero() -> None:
    rule = _override("campminder", person=1000002, household=1000001, session=1000102)
    plan = _plan([], overrides=[rule])
    assert _posts(plan) == {1: Decimal(0)}


def test_a_campminder_override_is_never_applied_where_campminders_money_cannot_be_placed() -> None:
    requests = [_request(), _request(LIAM, person=1000003, session=1000101)]
    unplaceable = _line("900", person=0, session=0)  # two sessions in the household: CampMinder can't tell which
    rule = _override("campminder", person=1000002, household=1000001, session=1000102)
    plan = _plan([], requests, [unplaceable], overrides=[rule])
    assert plan.creates == []
    assert "campminder_unmatched" in _kinds(plan)


def test_a_campminder_override_never_shares_a_family_and_session_the_sheet_loads() -> None:
    requests = [_request(), _request(LIAM, person=1000003)]
    rule = _override("campminder", person=1000003, household=1000001, session=1000102)
    plan = _plan([_award()], requests, [_line("3000")], overrides=[rule])
    assert set(_posts(plan, LIAM)) == set()
    assert "override_conflict" in _kinds(plan)


def test_a_campminder_override_on_a_withdrawn_request_is_reported_and_never_posted_onto() -> None:
    rule = _override("campminder", person=1000002, household=1000001, session=1000102)
    plan = _plan([], [_request(status="withdrawn")], overrides=[rule])
    assert plan.creates == []
    assert "request_not_live" in _kinds(plan)


def test_a_campminder_override_records_on_the_live_request_beside_its_duplicate_twin() -> None:
    rule = _override("campminder", person=1000002, household=1000001, session=1000102)
    plan = _plan([], [_request(LIAM, status="duplicate"), _request()], [_line("2500")], overrides=[rule])
    assert _posts(plan) == {1: Decimal(2500)}
    assert _posts(plan, LIAM) == {}
    assert "override_ambiguous" not in _kinds(plan)


def test_an_override_that_matches_nothing_is_reported_and_changes_nothing() -> None:
    rule = _override("campminder", person=1000077, household=1000076, session=1000102)
    plan = _plan([_award()], lines=[_line("3000")], overrides=[rule])
    assert _posts(plan) == {1: Decimal(3000)}
    (row,) = plan.report
    assert (row.kind, row.person_cm_id, row.household_cm_id, row.session_cm_id) == (
        "override_unused",
        1000077,
        1000076,
        1000102,
    )


def test_two_overrides_on_one_sheet_row_are_reported_and_neither_is_applied() -> None:
    award = _award(session_cm_id=1000101)
    rules = [
        _override("session", person=1000002, sheet_session=1000101, session=1000102),
        _override("skip", person=1000002, sheet_session=1000101),
    ]
    plan = _plan([award], lines=[_line("3000")], overrides=rules)
    assert plan.creates == []
    assert "override_ambiguous" in _kinds(plan)


@pytest.mark.parametrize(
    "fields",
    [
        {"decision": "session", "person": 1, "sheet_session": 2},  # which session?
        {"decision": "skip", "person": 1},  # which sheet session?
        {"decision": "person", "sheet_session": 2, "to_person": 3},  # which sheet row?
        {"decision": "campminder", "person": 1, "session": 2},  # which household?
        {"decision": "invent", "person": 1, "sheet_session": 2},
        {"decision": "skip", "person": 1, "sheet_session": 2, "amount": 500},  # an override never carries money
    ],
)
def test_an_override_missing_what_its_decision_needs_is_refused(fields: dict[str, Any]) -> None:
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        loader.LoadOverride.model_validate(fields)


# --- re-running ------------------------------------------------------------------------------------------------------


def _loaded(plan: loader.LoadPlan) -> list[DecisionEvent]:
    return _as_events(plan)


def test_a_second_run_with_the_same_sheet_and_ledger_writes_nothing() -> None:
    first = _plan([_award()], lines=[_line("3000")])
    second = _plan([_award()], lines=[_line("3000")], existing=_loaded(first))
    assert second.unchanged
    assert plan_writes(second, 2031) == []


def test_a_changed_run_replaces_exactly_the_loaders_rows() -> None:
    first = _plan([_award()], lines=[_line("3000")])
    staff = DecisionEvent(
        id="evstaff00000001", request_id=LIAM, round=1, kind="accept", created=T0, actor="registrar@example.com"
    )
    existing = [*_loaded(first), staff]
    requests = [_request(), _request(LIAM, person=1000003)]
    second = _plan([_award()], requests, [_line("3100")], existing=existing)
    writes = plan_writes(second, 2031)
    deleted = [w for w in writes if w.action == "delete"]
    assert sorted(w.record_id or "" for w in deleted) == sorted(e.id for e in _loaded(first))
    assert "evstaff00000001" not in {w.record_id for w in writes}
    created = [w for w in writes if w.action == "create"]
    bodies = [w.data for w in created if w.data is not None]
    assert [b["amount"] for b in bodies if b["event"] == "post"] == [3100.0]
    # Creates follow the deletes, and their ids sort in the order they must fold in.
    assert writes.index(created[0]) > writes.index(deleted[-1])
    ids = [w.record_id or "" for w in created]
    assert ids == sorted(ids)
    assert len(set(ids)) == len(ids)


def test_a_request_staff_have_decided_on_is_left_alone_and_reported() -> None:
    staff = DecisionEvent(
        id="evstaff00000001", request_id=EMMA, round=2, kind="ask", created=T0, amount=Decimal(500), actor="x@y.z"
    )
    plan = _plan([_award()], lines=[_line("3000")], existing=[staff])
    assert plan.creates == []
    assert _kinds(plan) == {"staff_rows": 1, "campminder_without_load": 1}


def test_a_loaded_request_staff_later_wrote_on_keeps_its_loaded_rows_on_a_re_run() -> None:
    requests = [_request(), _request(LIAM, person=1000003)]
    awards = [_award(), _award(row=3, person_cm_id=1000003)]
    first = _plan(awards, requests, [_line("3000"), _line("3000", txn=1000901, person=1000003)])
    staff = DecisionEvent(
        id="evstaff00000001", request_id=EMMA, round=1, kind="accept", created=T0, actor="registrar@example.com"
    )
    existing = [*_loaded(first), staff]
    second = _plan(awards, requests, [_line("3100"), _line("3000", txn=1000901, person=1000003)], existing=existing)
    deleted = {w.record_id for w in plan_writes(second, 2031) if w.action == "delete"}
    emmas = {e.id for e in existing if e.request_id == EMMA and e.actor == LOADER}
    assert emmas
    assert not deleted & emmas  # left as it is, as the staff_rows report says
    assert deleted == {e.id for e in existing if e.request_id == LIAM}
    assert "staff_rows" in _kinds(second)


# --- the season reads the loaded rows ----------------------------------------------------------------------------------


def test_the_budget_reads_the_loaded_rounds_as_posted_at_campminders_money() -> None:
    awards = [
        _award(stage="2 Accepted", appeal=Decimal(1000), r2=Decimal(600)),
        _award(
            row=3,
            person_cm_id=1000003,
            request=req(person_cm_id=1000003),
            stage="Full-cost program",
            extra=Decimal(1050),
            decision_type="full_cost_program",
        ),
    ]
    requests = [_request(), _request(LIAM, person=1000003)]
    lines = [_line("3400"), _line("4050", txn=1000901, person=1000003)]
    plan = _plan(awards, requests, lines, rules=OUTSIDE)
    rounds = fold_rounds(_as_events(plan))
    priced = {
        rid: price_request(
            RequestToPrice(
                request_id=rid,
                household_cm_id=1000001,
                live=True,
                application=app(),
                request=req(person_cm_id=person),
                blocked="",
                issues=(),
                rounds=rounds[rid],
                r1_ask=Decimal(4000),
            ),
            OUTSIDE,
        )
        for rid, person in ((EMMA, 1000002), (LIAM, 1000003))
    }
    total = season_budget(priced.values(), OUTSIDE, outside_grants={}).total.total
    assert total.posted == Decimal(3400)  # CampMinder's money; the full-cost round is outside the budget
    assert total.needs_offer == Decimal(0)
    assert total.committed == Decimal(3400)


# --- the workbook and the command --------------------------------------------------------------------------------------


_CONFIG = {
    "session_map": {"Session A": 1000102, "Session B": 1000101},
    "program_type_map": {"Summer": "summer"},
    "stage_decision_types": {"Full-cost program": "full_cost_program"},
    "unmatched_program": "other",
    "override_reason": "typed_household_total",
}


def _fill(ws: Any, columns: dict[str, tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    for key, (column, header) in columns.items():
        ws[f"{column}1"] = header.title()
        for number, values in enumerate(rows, start=2):
            ws[f"{column}{number}"] = values.get(key)


def _workbook(tmp_path: Path) -> tuple[Path, Path]:
    book = Workbook()
    book.remove(book.active)
    _fill(
        book.create_sheet("Raw Data"),
        RAW_COLUMNS,
        [
            {"unique_id": "U-1", "family_id": 1000001, "personal_id": 1000002, "py_gross": 60000, "cy_gross": 60000},
            {"unique_id": "U-2", "family_id": 1000001, "personal_id": 1000003, "py_gross": 60000, "cy_gross": 60000},
            {"unique_id": "U-3", "family_id": 1000004, "personal_id": None, "py_gross": 60000, "cy_gross": 60000},
        ],
    )
    _fill(
        book.create_sheet("Aid Calculator"),
        CALC_COLUMNS,
        [
            {
                "unique_id": "U-1",
                "include": "Yes",
                "stage": "2 Accepted",
                "session": "Session A",
                "ask": 4000,
                "r1": 3000,
                "appeal": 1000,
                "r2": 600,
                "total": 3600,
            },
            {
                "unique_id": "U-2",
                "include": "No",
                "stage": "Cancel After R1",
                "session": "Session A",
                "ask": 4000,
                "r1": 3000,
                "total": 3000,
            },
            {
                "unique_id": "U-3",
                "include": "Yes",
                "stage": "1 Accepted",
                "session": "Session A",
                "ask": 4000,
                "r1": 3000,
                "total": 3000,
            },
        ],
    )
    _fill(
        book.create_sheet("References"), REF_COLUMNS, [{"session": "Session A", "cost": 4000, "program_type": "Summer"}]
    )
    _fill(book.create_sheet("Grants"), GRANT_COLUMNS, [])
    path = tmp_path / "fictional.xlsx"
    book.save(path)
    config = tmp_path / "config.json"
    config.write_text(json.dumps(_CONFIG))
    return path, config


def test_the_workbook_gives_one_award_per_included_row_and_the_sheets_own_totals(tmp_path: Path) -> None:
    path, config = _workbook(tmp_path)
    awards, totals = loader.read_awards(path, ParityConfig.model_validate_json(config.read_text()), RULES)
    assert [(a.row, a.person_cm_id, a.household_cm_id, a.session_cm_id) for a in awards] == [
        (2, 1000002, 1000001, 1000102),
        (4, None, 1000004, 1000102),
    ]
    assert (awards[0].appeal, awards[0].r2, awards[0].stage) == (Decimal(1000), Decimal(600), "2 Accepted")
    assert (totals.included, totals.excluded) == (2, 1)
    assert (totals.r1, totals.r2, totals.r3) == (Decimal(6000), Decimal(600), Decimal(0))


class _Store:
    def __init__(self, existing: list[DecisionEvent] | None = None) -> None:
        self.events = existing or []
        self.commits: list[list[AidWrite]] = []

    async def rules(self, year: int) -> tuple[AidRules, int] | None:
        return RULES, 3

    async def requests(self, year: int) -> list[LoadRequest]:
        return [_request()]

    async def events_of(self, year: int) -> list[DecisionEvent]:
        return self.events

    async def camp_lines(self, year: int) -> list[CampLine]:
        return [_line("3400")]

    async def commit(self, writes: Sequence[AidWrite]) -> None:
        self.commits.append(list(writes))


def _run(tmp_path: Path, store: _Store, *extra: str, monkeypatch: pytest.MonkeyPatch) -> int:
    path, config = _workbook(tmp_path)
    monkeypatch.setattr(loader, "open_store", lambda: store)
    return loader.main(
        [
            "--workbook",
            str(path),
            "--config",
            str(config),
            "--year",
            "2031",
            "--report",
            str(tmp_path / "r.csv"),
            *extra,
        ]
    )


def test_a_dry_run_prints_aggregates_writes_the_report_and_writes_nothing(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    store = _Store()
    assert _run(tmp_path, store, monkeypatch=monkeypatch) == 0
    out = capsys.readouterr().out
    assert store.commits == []
    assert "Dry run" in out
    assert "3,600" in out  # the sheet's rounds and CampMinder's money, as totals
    assert "3,400" in out
    assert "1000002" not in out  # ids reach the report file, never the console
    assert "1000001" not in out
    report = (tmp_path / "r.csv").read_text()
    assert "campminder_differs" in report
    assert "no_campminder_id" in report


def test_write_refuses_without_a_pocketbase_url(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("POCKETBASE_URL", raising=False)
    store = _Store()
    assert _run(tmp_path, store, "--write", monkeypatch=monkeypatch) == 2
    assert store.commits == []


def test_write_commits_once_and_a_second_run_finds_it_already_loaded(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("POCKETBASE_URL", "http://pocketbase.invalid:8090")
    store = _Store()
    assert _run(tmp_path, store, "--write", monkeypatch=monkeypatch) == 0
    (writes,) = store.commits
    created = [w for w in writes if w.action == "create"]
    store.events = [
        DecisionEvent(
            id=w.record_id or "",
            request_id=w.data["request"],
            round=w.data["round"],
            kind=w.data["event"],
            created=T0,
            amount=Decimal(str(w.data["amount"])) if "amount" in w.data else None,
            decision_type=w.data.get("decision_type", ""),
            lock_source=w.data.get("lock_source", ""),
            rules_version=w.data.get("rules_version"),
            snapshot=w.data.get("snapshot"),
            note=w.data.get("note", ""),
            actor=w.data["actor"],
        )
        for w in created
        if w.data is not None
    ]
    capsys.readouterr()
    assert _run(tmp_path, store, "--write", monkeypatch=monkeypatch) == 0
    assert len(store.commits) == 1
    assert "already loaded" in capsys.readouterr().out


def test_the_command_reads_the_overrides_file(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    overrides = tmp_path / "overrides.json"
    overrides.write_text(
        json.dumps([{"decision": "skip", "person": 1000002, "sheet_session": 1000102, "note": "fictional"}])
    )
    store = _Store()
    assert _run(tmp_path, store, "--overrides", str(overrides), monkeypatch=monkeypatch) == 0
    assert "override_skip" in (tmp_path / "r.csv").read_text()


def test_the_command_refuses_an_invalid_overrides_file_without_echoing_its_ids(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    overrides = tmp_path / "overrides.json"
    overrides.write_text(json.dumps([{"decision": "session", "person": 1000002, "sheet_session": 1000102}]))
    store = _Store()
    assert _run(tmp_path, store, "--overrides", str(overrides), monkeypatch=monkeypatch) == 2
    err = capsys.readouterr().err
    assert "Nothing was written" in err
    assert "1000002" not in err
    assert store.commits == []


# --- --sheet-id: the live sheet through the service account (the owner's January reload path) -------------------------


def _credentials(tmp_path: Path) -> Path:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    pem = key.private_bytes(
        serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()
    ).decode()
    path = tmp_path / "service-account.json"
    path.write_text(
        json.dumps(
            {
                "type": "service_account",
                "client_email": "loader@fictional-project.iam.gserviceaccount.com",
                "private_key": pem,
                "token_uri": "https://oauth2.invalid/token",
            }
        )
    )
    return path


class _Http:
    def __init__(self, content: bytes) -> None:
        self.content = content
        self.calls: list[tuple[str, str, dict[str, Any]]] = []

    def post(self, url: str, **kwargs: Any) -> Any:
        self.calls.append(("POST", url, kwargs))
        return _Response(json_body={"access_token": "token-1", "token_type": "Bearer"})

    def get(self, url: str, **kwargs: Any) -> Any:
        self.calls.append(("GET", url, kwargs))
        return _Response(content=self.content)


class _Response:
    def __init__(self, *, json_body: dict[str, Any] | None = None, content: bytes = b"") -> None:
        self._json, self.content = json_body, content

    def raise_for_status(self) -> None:
        return None

    def json(self) -> dict[str, Any]:
        assert self._json is not None
        return self._json


def test_a_sheet_id_is_exported_as_xlsx_read_only_through_the_service_account(tmp_path: Path) -> None:
    http = _Http(b"PK-fictional-xlsx")
    assert loader.export_sheet("sheet-1", _credentials(tmp_path), http=http) == b"PK-fictional-xlsx"
    (_, token_url, token), (_, export_url, export) = http.calls
    assert token_url == "https://oauth2.invalid/token"
    assert token["data"]["grant_type"] == "urn:ietf:params:oauth:grant-type:jwt-bearer"
    import jwt

    claims = jwt.decode(token["data"]["assertion"], options={"verify_signature": False})
    assert claims["scope"] == "https://www.googleapis.com/auth/drive.readonly"  # never a write scope
    assert export_url == "https://www.googleapis.com/drive/v3/files/sheet-1/export"
    assert export["params"]["mimeType"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert export["headers"]["Authorization"] == "Bearer token-1"


def test_the_command_reads_the_exported_sheet_in_memory(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], monkeypatch: pytest.MonkeyPatch
) -> None:
    path, config = _workbook(tmp_path)
    monkeypatch.setattr(loader, "export_sheet", lambda sheet_id, credentials, http=None: path.read_bytes())
    monkeypatch.setattr(loader, "open_store", lambda: _Store())
    before = set(tmp_path.iterdir())
    code = loader.main(
        ["--sheet-id", "sheet-1", "--credentials", str(tmp_path / "unused.json"), "--config", str(config),
         "--year", "2031", "--report", str(tmp_path / "r.csv")]
    )  # fmt: skip
    assert code == 0
    assert "Dry run" in capsys.readouterr().out
    assert set(tmp_path.iterdir()) - before == {tmp_path / "r.csv"}  # the sheet is never written to disk


def test_workbook_and_sheet_id_are_one_or_the_other(tmp_path: Path) -> None:
    with pytest.raises(SystemExit):
        loader.main(["--workbook", "a.xlsx", "--sheet-id", "s", "--config", "c", "--report", "r"])
    with pytest.raises(SystemExit):
        loader.main(["--config", "c", "--report", "r"])
