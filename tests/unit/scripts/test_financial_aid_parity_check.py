"""The parity harness on a three-row synthetic workbook built from fictional rules.

Row 2 is a summer camper, row 3 a family-camp household with a typed cost, row 4
a session the catalog does not know. Each cached sheet value is what the
fictional rules give, so sheet mode matches every row; a Grants-tab grant for
row 2's camper makes correct mode differ on that row only.
"""

import json
import re
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from scripts.financial_aid.parity_check import CALC_COLUMNS, GRANT_COLUMNS, RAW_COLUMNS, REF_COLUMNS, main
from tests.unit.bunking.financial_aid.fixtures import fictional_rules_json

_CONFIG = {
    "session_map": {"Session A": 1000102, "Family Weekend": 1000201},
    "program_type_map": {"Summer": "summer", "Family": "family_camp"},
    "unmatched_program": "no_match",
    "override_reason": "typed_household_total",
}


def _fill(ws: Any, columns: dict[str, tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    for key, (column, header) in columns.items():
        ws[f"{column}1"] = header.title()  # headers compare case-insensitively
        for number, values in enumerate(rows, start=2):
            ws[f"{column}{number}"] = values.get(key)


def _write_workbook(path: Path, *, r1_row_2: int = 3000, raw_header_n: str | None = None) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    raw = workbook.create_sheet("Raw Data")
    _fill(
        raw,
        RAW_COLUMNS,
        [
            {"unique_id": "U-1", "personal_id": 1000002, "py_gross": 60000, "cy_gross": 60000, "bipoc": "No"},
            {"unique_id": "U-2", "personal_id": 1000004, "py_gross": 100000, "cy_gross": 100000},
            {"unique_id": "U-3", "personal_id": 1000005, "py_gross": 250000, "cy_gross": 250000},
        ],
    )
    if raw_header_n is not None:
        raw["N1"] = raw_header_n
    calc = workbook.create_sheet("Aid Calculator")
    _fill(
        calc,
        CALC_COLUMNS,
        [
            {
                "unique_id": "U-1",
                "stage": "Accepted",
                "session": "Session A",
                "ask": 4000,
                "income": 60000,
                "income_tier": 2,
                "final_tier": 2,
                "potential": 3000,
                "r1": r1_row_2,
                "appeal": 1000,
                "r2": 600,
                "total": 3600,
            },
            {
                "unique_id": "U-2",
                "stage": "Accepted",
                "session": "Family Weekend",
                "ask": 4000,
                "override": 2100,
                "income": 100000,
                "income_tier": 3,
                "final_tier": 3,
                "potential": 1155,
                "r1": 1155,
                "total": 1155,
            },
            {
                "unique_id": "U-3",
                "stage": "Accepted",
                "session": "Mystery Week",
                "ask": 4000,
                "income": 250000,
                "income_tier": 6,
                "final_tier": 6,
                "potential": 100,
                "r1": 100,
                "total": 100,
            },
        ],
    )
    _fill(
        workbook.create_sheet("References"),
        REF_COLUMNS,
        [
            {"session": "Session A", "cost": 4000, "program_type": "Summer"},
            {"session": "Family Weekend", "cost": None, "program_type": "Family"},
        ],
    )
    _fill(
        workbook.create_sheet("Grants"),
        GRANT_COLUMNS,
        [
            {"cancelled": None, "camper_id": 1000002, "amount": 500},
            {"cancelled": "Yes", "camper_id": 1000004, "amount": 900},
        ],
    )
    workbook.save(path)


def _files(
    tmp_path: Path, *, rules_changes: dict[str, Any] | None = None, config: dict[str, Any] | None = None
) -> list[str]:
    rules = fictional_rules_json()
    rules["programs"]["no_match"] = {
        "label": "No session match",
        "session_cm_ids": [],
        "r1_table": None,
        "r2_table": None,
        "equity_class": None,
        "budget_pool": None,
        "cost_source": "catalog",
    }
    for key, value in (rules_changes or {}).items():
        rules[key] = value
    (tmp_path / "rules.json").write_text(json.dumps(rules))
    (tmp_path / "config.json").write_text(json.dumps(config or _CONFIG))
    return [
        "--rules",
        str(tmp_path / "rules.json"),
        "--config",
        str(tmp_path / "config.json"),
        "--sheet",
        str(tmp_path / "sheet.xlsx"),
    ]


def test_a_matching_sheet_passes_every_column(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _write_workbook(tmp_path / "sheet.xlsx")
    assert main([*_files(tmp_path), "--mode", "sheet"]) == 0
    out = capsys.readouterr().out
    for label in (
        "adjusted income (K)",
        "final tier (M)",
        "R1 potential (N)",
        "Round 1 (Q)",
        "Round 2 (W)",
        "total (J)",
    ):
        assert re.search(re.escape(label) + r"\s+3/3", out), label
    assert "RESULT sheet-grant: PASS" in out


def test_one_wrong_cell_fails_and_names_only_its_row(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _write_workbook(tmp_path / "sheet.xlsx", r1_row_2=2999)
    assert main([*_files(tmp_path), "--mode", "sheet"]) == 1
    out = capsys.readouterr().out
    assert re.search(r"Round 1 \(Q\)\s+2/3", out)
    assert "mismatched sheet rows: 2" in out
    assert "RESULT sheet-grant: FAIL" in out


def test_correct_grant_mode_reports_differences_without_failing(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _write_workbook(tmp_path / "sheet.xlsx")
    assert main([*_files(tmp_path), "--mode", "both"]) == 0
    out = capsys.readouterr().out
    correct = out[out.index("mode correct-grant") :]
    assert re.search(r"R1 potential \(N\)\s+2/3", correct)
    assert re.search(r"adjusted income \(K\)\s+3/3", correct)


def test_the_output_names_no_person_and_no_row_id(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _write_workbook(tmp_path / "sheet.xlsx")
    main([*_files(tmp_path), "--mode", "both"])
    out = capsys.readouterr().out
    for private in ("1000002", "1000004", "U-1", "U-2"):
        assert private not in out


def test_a_moved_column_is_a_layout_error(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _write_workbook(tmp_path / "sheet.xlsx", raw_header_n="Something Else")
    assert main([*_files(tmp_path), "--mode", "sheet"]) == 2
    assert "RESULT: LAYOUT" in capsys.readouterr().out


def test_a_catalog_session_missing_from_the_map_fails(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _write_workbook(tmp_path / "sheet.xlsx")
    config = {**_CONFIG, "session_map": {"Family Weekend": 1000201}}
    assert main([*_files(tmp_path, config=config), "--mode", "sheet"]) == 1
    assert "session a" in capsys.readouterr().out


def test_rules_with_errors_stop_before_the_sheet_is_read(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    _write_workbook(tmp_path / "sheet.xlsx")
    pools = {
        "camp_pool": {"label": "Camp", "share_pct": "79"},
        "weekend_pool": {"label": "W", "share_pct": "15"},
        "bmitzvah_pool": {"label": "B", "share_pct": "5"},
    }
    budget = {**fictional_rules_json()["budget"], "pools": pools}
    assert main([*_files(tmp_path, rules_changes={"budget": budget}), "--mode", "sheet"]) == 1
    assert "validation errors" in capsys.readouterr().out
